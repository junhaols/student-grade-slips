#!/usr/bin/env python3
"""Turn a one/two-header-row grade workbook into cuttable student slips."""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.properties import PageSetupProperties


EMPTY_MARKERS = {"", "-", "—", "–"}


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def meaningful(value: object) -> bool:
    return value is not None and str(value).strip() not in EMPTY_MARKERS


def display(value: object) -> str:
    return "" if value is None else str(value)


def load_source(path: Path, sheet_name: str | None, header_rows: int, name_label: str,
                drop_labels: list[str]):
    values_book = load_workbook(path, data_only=True, read_only=False)
    formulas_book = load_workbook(path, data_only=False, read_only=False)
    if sheet_name and sheet_name not in values_book.sheetnames:
        raise ValueError(f"工作表不存在：{sheet_name}; 可用：{values_book.sheetnames}")
    if not sheet_name and len(values_book.sheetnames) != 1:
        raise ValueError(f"有多个工作表，请指定 --sheet：{values_book.sheetnames}")
    sheet_name = sheet_name or values_book.sheetnames[0]
    ws = values_book[sheet_name]
    formula_ws = formulas_book[sheet_name]
    if ws.max_row <= header_rows:
        raise ValueError("表头后没有数据行")

    headers = []
    parent = ""
    for col in range(1, ws.max_column + 1):
        top = ws.cell(1, col).value
        if meaningful(top):
            parent = display(top).strip()
        sub = display(ws.cell(2, col).value).strip() if header_rows == 2 else ""
        if not parent and not sub:
            raise ValueError(f"第 {get_column_letter(col)} 列没有表头，无法安全识别")
        headers.append({"col": col, "parent": parent or sub, "sub": sub})

    def matches(label: str) -> list[int]:
        return [h["col"] for h in headers if label == h["parent"] or
                label == f'{h["parent"]}/{h["sub"]}']

    name_cols = matches(name_label)
    if len(name_cols) != 1:
        raise ValueError(f"姓名列匹配到 {len(name_cols)} 列，请用唯一的 --name-column 指定")
    name_col = name_cols[0]

    rows = []
    for rownum in range(header_rows + 1, ws.max_row + 1):
        row = tuple(ws.cell(rownum, col).value for col in range(1, ws.max_column + 1))
        if not any(meaningful(v) for v in row):
            continue
        if not meaningful(row[name_col - 1]):
            raise ValueError(f"源表第 {rownum} 行有数据但姓名为空")
        for col, value in enumerate(row, 1):
            if formula_ws.cell(rownum, col).data_type == "f" and value is None:
                raise ValueError(f"源表 {get_column_letter(col)}{rownum} 的公式无缓存结果")
        rows.append((rownum, row))
    if not rows:
        raise ValueError("没有有效学生行")

    drop_cols = set()
    for label in drop_labels:
        found = matches(label)
        if not found:
            raise ValueError(f"待删除列未找到：{label}")
        drop_cols.update(found)
    if name_col in drop_cols:
        raise ValueError("姓名列不能删除")
    empty_cols = {col for col in range(1, ws.max_column + 1)
                  if all(not meaningful(row[col - 1]) for _, row in rows)}
    kept = [col for col in range(1, ws.max_column + 1)
            if col not in drop_cols and col not in empty_cols]
    if not kept:
        raise ValueError("没有可保留的列")
    return sheet_name, headers, rows, kept, sorted(drop_cols), sorted(empty_cols), name_col


def make_excel(path: Path, headers: list[dict], rows: list[tuple], kept: list[int],
               name_col: int, header_rows: int, per_page: int) -> None:
    book = Workbook()
    ws = book.active
    ws.title = "可裁剪成绩单"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True, autoPageBreaks=False)
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins = PageMargins(left=.18, right=.18, top=.20, bottom=.20,
                                  header=0, footer=0)
    ws.print_options.horizontalCentered = True
    ws.sheet_view.zoomScale = 85
    for j, col in enumerate(kept, 1):
        width = 12 if col == name_col else (5.5 if headers[col - 1]["sub"] in
                                            ("成绩", "等级分") else 4.5)
        ws.column_dimensions[get_column_letter(j)].width = width

    thin = Side(style="thin", color="475569")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    top_fill = PatternFill("solid", fgColor="DDEBF3")
    sub_fill = PatternFill("solid", fgColor="EEF5F9")
    header_font = Font(name="PingFang SC", size=8, bold=True, color="15283B")
    data_font = Font(name="Arial Narrow", size=8.5, color="15283B")
    for i, (_, record) in enumerate(rows):
        first = i * (header_rows + 2) + 1
        data_row = first + header_rows
        spacer = data_row + 1
        for rownum in range(first, data_row + 1):
            ws.row_dimensions[rownum].height = 36 if rownum == first else (26 if rownum < data_row else 23)
            for j in range(1, len(kept) + 1):
                cell = ws.cell(rownum, j)
                cell.border = border
                cell.alignment = Alignment(horizontal="center", vertical="center",
                                           wrap_text=rownum < data_row)
                cell.font = header_font if rownum < data_row else data_font
                if rownum == first:
                    cell.fill = top_fill
                elif rownum < data_row:
                    cell.fill = sub_fill
        ws.row_dimensions[spacer].height = 16
        ws.cell(spacer, 1).value = ""

        if header_rows == 2:
            j = 1
            while j <= len(kept):
                parent = headers[kept[j - 1] - 1]["parent"]
                end = j
                while end < len(kept) and headers[kept[end] - 1]["parent"] == parent:
                    end += 1
                if end > j:
                    ws.merge_cells(start_row=first, start_column=j,
                                   end_row=first, end_column=end)
                if end == j and not headers[kept[j - 1] - 1]["sub"]:
                    ws.merge_cells(start_row=first, start_column=j,
                                   end_row=first + 1, end_column=j)
                ws.cell(first, j).value = parent
                for x in range(j, end + 1):
                    sub = headers[kept[x - 1] - 1]["sub"]
                    if sub:
                        ws.cell(first + 1, x).value = sub.replace("排名", "\n排名") if sub.endswith("排名") else sub
                j = end + 1
        else:
            for j, col in enumerate(kept, 1):
                ws.cell(first, j).value = headers[col - 1]["parent"]
        for j, col in enumerate(kept, 1):
            cell = ws.cell(data_row, j)
            cell.value = record[col - 1]
            if col == name_col:
                cell.font = Font(name="PingFang SC", size=8.5, bold=True, color="15283B")
        if (i + 1) % per_page == 0 and i + 1 < len(rows):
            ws.row_breaks.append(Break(id=spacer))
    ws.print_area = f"A1:{get_column_letter(len(kept))}{len(rows) * (header_rows + 2)}"
    book.save(path)


def make_word(path: Path, headers: list[dict], rows: list[tuple], kept: list[int],
              name_col: int, header_rows: int) -> None:
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt

    def put(cell, value, *, bold=False):
        cell.text = ""
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(display(value))
        run.bold = bold
        run.font.name = "Arial Narrow"
        run.font.size = Pt(5.8 if len(kept) > 35 else 6.5)
        fonts = run._element.get_or_add_rPr().rFonts
        if fonts is not None:
            fonts.set(qn("w:eastAsia"), "PingFang SC")

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(297), Mm(210)
    section.top_margin = section.bottom_margin = Mm(7)
    section.left_margin = section.right_margin = Mm(7)
    section.header_distance = section.footer_distance = Mm(3)
    weights = [1.8 if col == name_col else 1.1 if headers[col - 1]["sub"] in
               ("成绩", "等级分") else 1 for col in kept]
    widths = [283 * weight / sum(weights) for weight in weights]
    for index, (_, record) in enumerate(rows):
        table = doc.add_table(rows=header_rows + 1, cols=len(kept))
        table.autofit = False
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"
        for row_index, row in enumerate(table.rows):
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            row.height = Mm(8 if row_index == 0 else 7)
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            for cell, width in zip(row.cells, widths):
                cell.width = Mm(width)
        for col, width in zip(table.columns, widths):
            col.width = Mm(width)
        if header_rows == 2:
            j = 0
            while j < len(kept):
                parent = headers[kept[j] - 1]["parent"]
                end = j + 1
                while end < len(kept) and headers[kept[end] - 1]["parent"] == parent:
                    end += 1
                if end - j > 1:
                    put(table.cell(0, j).merge(table.cell(0, end - 1)), parent, bold=True)
                elif not headers[kept[j] - 1]["sub"]:
                    put(table.cell(0, j).merge(table.cell(1, j)), parent, bold=True)
                else:
                    put(table.cell(0, j), parent, bold=True)
                for x in range(j, end):
                    sub = headers[kept[x] - 1]["sub"]
                    if sub:
                        put(table.cell(1, x), sub, bold=True)
                j = end
        else:
            for j, col in enumerate(kept):
                put(table.cell(0, j), headers[col - 1]["parent"], bold=True)
        for j, col in enumerate(kept):
            put(table.cell(header_rows, j), record[col - 1], bold=col == name_col)
        for row in table.rows[:-1]:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True
        spacer = doc.add_paragraph(" ")
        spacer.paragraph_format.space_after = Mm(4)
        spacer.runs[0].font.size = Pt(1)
    doc.save(path)


def verify_excel(path: Path, headers: list[dict], rows: list[tuple], kept: list[int],
                 header_rows: int) -> None:
    ws = load_workbook(path, data_only=True).active
    if ws.page_setup.orientation != "landscape" or str(ws.page_setup.paperSize) != str(ws.PAPERSIZE_A4):
        raise AssertionError("Excel 页面设置不是 A4 横版")
    for i, (_, record) in enumerate(rows):
        first = i * (header_rows + 2) + 1
        for j, col in enumerate(kept, 1):
            actual = ws.cell(first + header_rows, j).value
            expected = record[col - 1]
            if type(actual) is not type(expected) or actual != expected:
                raise AssertionError(f"Excel 第 {i+1} 名学生、第 {j} 列与源表不符")
            header = headers[col - 1]
            if header_rows == 1:
                actual_parent = ws.cell(first, j).value
            else:
                actual_parent = ws.cell(first, j).value
                if actual_parent is None:
                    for left in range(j - 1, 0, -1):
                        actual_parent = ws.cell(first, left).value
                        if actual_parent is not None:
                            break
                actual_sub = ws.cell(first + 1, j).value
                expected_sub = (header["sub"].replace("排名", "\n排名")
                                if header["sub"].endswith("排名") else header["sub"])
                if header["sub"] and actual_sub != expected_sub:
                    raise AssertionError(f"Excel 第 {i+1} 名学生、第 {j} 列子表头不符")
            if actual_parent != header["parent"]:
                raise AssertionError(f"Excel 第 {i+1} 名学生、第 {j} 列主表头不符")
    if ws.max_column != len(kept) or ws.max_row != len(rows) * (header_rows + 2):
        raise AssertionError("Excel 行列数不符")


def verify_word(path: Path, headers: list[dict], rows: list[tuple], kept: list[int],
                header_rows: int) -> None:
    from docx import Document
    from docx.shared import Mm
    doc = Document(path)
    if len(doc.tables) != len(rows):
        raise AssertionError("Word 成绩条数量不符")
    section = doc.sections[0]
    if abs(section.page_width - Mm(297)) > 1000 or abs(section.page_height - Mm(210)) > 1000:
        raise AssertionError("Word 页面设置不是 A4 横版")
    for i, (_, record) in enumerate(rows):
        table = doc.tables[i]
        actual = [cell.text for cell in table.rows[header_rows].cells]
        expected = [display(record[col - 1]) for col in kept]
        if actual != expected:
            raise AssertionError(f"Word 第 {i+1} 名学生的成绩与源表不符")
        for j, col in enumerate(kept):
            header = headers[col - 1]
            if table.cell(0, j).text != header["parent"]:
                raise AssertionError(f"Word 第 {i+1} 名学生、第 {j+1} 列主表头不符")
            if header_rows == 2 and header["sub"] and table.cell(1, j).text != header["sub"]:
                raise AssertionError(f"Word 第 {i+1} 名学生、第 {j+1} 列子表头不符")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--sheet")
    parser.add_argument("--header-rows", type=int, choices=(1, 2), default=2)
    parser.add_argument("--name-column", default="姓名")
    parser.add_argument("--drop-columns", nargs="*", default=[])
    parser.add_argument("--format", choices=("excel", "word", "both"), default="both")
    parser.add_argument("--rows-per-page", type=int, default=5,
                        help="Excel 每页成绩条数；Word 根据文字长度自然分页")
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    if args.rows_per_page < 1:
        parser.error("--rows-per-page 必须大于 0")
    if not args.input.is_file():
        parser.error(f"源文件不存在：{args.input}")
    args.output_dir.mkdir(parents=True, exist_ok=True)
    sheet_name, headers, rows, kept, dropped, empty, name_col = load_source(
        args.input, args.sheet, args.header_rows, args.name_column, args.drop_columns)
    stem = args.input.stem + "_可裁剪成绩单"
    kinds = ("excel", "word") if args.format == "both" else (args.format,)
    planned = [args.output_dir / (stem + (".xlsx" if kind == "excel" else ".docx"))
               for kind in kinds] + [args.output_dir / "核验报告.json"]
    for path in planned:
        if path.exists():
            raise FileExistsError(f"输出文件已存在，请使用新的输出目录：{path}")
    artifacts = {}
    for kind, suffix in (("excel", ".xlsx"), ("word", ".docx")):
        if args.format not in (kind, "both"):
            continue
        path = args.output_dir / (stem + suffix)
        if kind == "excel":
            make_excel(path, headers, rows, kept, name_col, args.header_rows, args.rows_per_page)
            verify_excel(path, headers, rows, kept, args.header_rows)
        else:
            make_word(path, headers, rows, kept, name_col, args.header_rows)
            verify_word(path, headers, rows, kept, args.header_rows)
        artifacts[kind] = {"path": str(path), "sha256": digest(path)}
    report = {
        "verified": True,
        "source": str(args.input),
        "source_sha256": digest(args.input),
        "sheet": sheet_name,
        "header_rows": args.header_rows,
        "source_row_numbers": [rownum for rownum, _ in rows],
        "student_count": len(rows),
        "name_source_column": name_col,
        "kept_source_columns": kept,
        "explicitly_dropped_columns": dropped,
        "empty_columns": empty,
        "compared_data_cells_per_artifact": len(rows) * len(kept),
        "excel_rows_per_page": args.rows_per_page,
        "artifacts": artifacts,
    }
    report_path = args.output_dir / "核验报告.json"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
